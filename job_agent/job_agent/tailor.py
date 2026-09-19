from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .llm import Ollama, generate_materials


def tailor_to_docx(cfg, profile, job):
    llm_cfg = cfg["llm"]
    llm = Ollama(llm_cfg["base_url"], llm_cfg["model"], llm_cfg.get("temperature", 0.15))
    data = generate_materials(llm, profile, job)
    out = Path(cfg["report_dir"]) / f"{job.company}_{job.id.replace(':', '_')}_application"
    out.mkdir(parents=True, exist_ok=True)
    (out / "cover_letter.txt").write_text(data.get("cover_letter", ""), encoding="utf-8")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    doc.add_heading(profile["name"], 0)
    doc.add_paragraph(data.get("summary", ""))
    doc.add_heading("Experience", 1)
    for b in data.get("cv_bullets", []):
        doc.add_paragraph(b, style="List Bullet")
    doc.save(out / "tailored_cv.docx")
    return out
